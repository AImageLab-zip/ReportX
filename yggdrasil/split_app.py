"""
Annotation Split Manager — cross-platform desktop app
Run with: .venv/bin/python yggrasil/split_app.py
"""

import json
import math
import random
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
import pandas as pd
from PyQt6.QtCore import (
    Qt, QAbstractTableModel, QModelIndex, QVariant, QSize, QProcess,
)
from PyQt6.QtGui import QColor, QFont, QIcon
from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QLabel, QSpinBox, QLineEdit, QPushButton, QTabWidget, QTableWidget,
    QTableWidgetItem, QHeaderView, QGroupBox, QFormLayout, QMessageBox,
    QFileDialog, QScrollArea, QTableView, QAbstractItemView,
)

INPUT_JSON = Path(__file__).parent.parent / "split_files" / "BraTS23.json"
DEFAULT_EXCEL = Path(__file__).parent / "annotation_splits.xlsx"
REPORTX_JSON = Path(__file__).parent / "reportx.json"
REPORTX_KEY = "REPORTX"

# ── Dark Catppuccin palette ─────────────────────────────────────────────────
BG        = "#1e1e2e"
SURFACE   = "#313244"
OVERLAY   = "#45475a"
TEXT      = "#cdd6f4"
SUBTEXT   = "#a6adc8"
ACCENT    = "#89b4fa"
GREEN     = "#a6e3a1"
RED       = "#f38ba8"
YELLOW    = "#f9e2af"
BLUE      = "#89b4fa"

MAIN_STYLE = f"""
QMainWindow, QWidget {{ background-color: {BG}; color: {TEXT}; }}
QLabel {{ color: {TEXT}; }}
QGroupBox {{
    border: 1px solid {OVERLAY};
    border-radius: 6px;
    margin-top: 8px;
    font-weight: bold;
    color: {ACCENT};
}}
QGroupBox::title {{ subcontrol-origin: margin; left: 10px; padding: 0 4px; }}
"""

BUTTON_STYLE = f"""
QPushButton {{
    background-color: {SURFACE};
    color: {TEXT};
    border: 1px solid {OVERLAY};
    border-radius: 6px;
    padding: 8px 16px;
    font-size: 13px;
    font-weight: 500;
}}
QPushButton:hover {{ background-color: {OVERLAY}; border-color: {ACCENT}; }}
QPushButton:pressed {{ background-color: {ACCENT}; color: #1e1e2e; }}
"""

PRIMARY_BTN_STYLE = f"""
QPushButton {{
    background-color: {BLUE};
    color: #1e1e2e;
    border: none;
    border-radius: 6px;
    padding: 8px 16px;
    font-size: 13px;
    font-weight: bold;
}}
QPushButton:hover {{ background-color: #a6d4fa; }}
"""

ADD_BTN_STYLE = f"""
QPushButton {{
    background-color: {GREEN};
    color: #1e1e2e;
    border: none;
    border-radius: 6px;
    padding: 6px 14px;
    font-size: 13px;
    font-weight: bold;
}}
QPushButton:hover {{ background-color: #b9f0b4; }}
"""

DEL_BTN_STYLE = f"""
QPushButton {{
    background-color: {RED};
    color: #1e1e2e;
    border: none;
    border-radius: 6px;
    padding: 6px 14px;
    font-size: 13px;
    font-weight: bold;
}}
QPushButton:hover {{ background-color: #f5a0b5; }}
"""

SPIN_STYLE = f"""
QSpinBox {{
    background-color: {SURFACE};
    color: {TEXT};
    border: 1px solid {OVERLAY};
    border-radius: 4px;
    padding: 4px 6px;
}}
"""

LINE_EDIT_STYLE = f"""
QLineEdit {{
    background-color: {SURFACE};
    color: {TEXT};
    border: 1px solid {OVERLAY};
    border-radius: 4px;
    padding: 4px 6px;
}}
QLineEdit:focus {{ border: 1px solid {ACCENT}; }}
"""

TAB_STYLE = f"""
QTabWidget::pane {{ border: 1px solid {OVERLAY}; background: {BG}; }}
QTabBar::tab {{
    background: {SURFACE};
    color: {SUBTEXT};
    padding: 8px 16px;
    margin-right: 2px;
    border-top-left-radius: 6px;
    border-top-right-radius: 6px;
}}
QTabBar::tab:selected {{ background: {OVERLAY}; color: {ACCENT}; font-weight: bold; }}
QTabBar::tab:hover {{ color: {ACCENT}; }}
"""

TABLE_STYLE = f"""
QTableWidget {{
    background-color: {BG};
    color: {TEXT};
    gridline-color: {OVERLAY};
    border: none;
    font-size: 12px;
}}
QHeaderView::section {{
    background-color: {SURFACE};
    color: {ACCENT};
    padding: 6px;
    border: 1px solid {OVERLAY};
    font-weight: bold;
}}
QTableWidget::item:selected {{ background-color: {ACCENT}; color: #1e1e2e; }}
QTableWidget::item {{ padding: 2px; }}
"""


# ── Split logic ─────────────────────────────────────────────────────────────

def patient_id(sample: str) -> str:
    """BraTS-GLI-00002-000 → BraTS-GLI-00002"""
    return "-".join(sample.split("-")[:-1])


def load_reportx_samples() -> list:
    """Cases always pre-assigned to the REPORTX clinician (one fulfilled repeat each)."""
    try:
        with open(REPORTX_JSON) as f:
            samples = json.load(f)
    except FileNotFoundError:
        return []
    return sorted(set(samples))


def build_pool(data: dict, train_r: int, val_r: int, test_r: int, reportx_ids: set = frozenset()) -> list:
    """Build the assignment pool, accounting for one repeat already fulfilled by REPORTX."""
    pool = []
    repeats = {"train": train_r, "val": val_r, "test": test_r}
    for subset, base_repeat in repeats.items():
        for s in data[subset]:
            n = base_repeat - (1 if s in reportx_ids else 0)
            if n > 0:
                pool.extend([s] * n)
    return pool


def spread_pool(pool: list, rng: random.Random) -> list:
    by_patient = defaultdict(list)
    for s in pool:
        by_patient[patient_id(s)].append(s)
    groups = sorted(by_patient.values(), key=len, reverse=True)
    interleaved = []
    while any(groups):
        rng.shuffle(groups)
        for g in groups:
            if g:
                interleaved.append(g.pop())
        groups = [g for g in groups if g]
    return interleaved


def assign_to_buckets(pool: list, n_clinicians: int, split_size: int, prefix: str, seed: int):
    """Distribute pool across n_clinicians buckets of split_size each. Overflow → REMAINING."""
    rng = random.Random(seed)
    pool = spread_pool(pool, rng)

    keys = [f"{prefix}_{i+1:02d}" for i in range(n_clinicians)]
    buckets = {k: [] for k in keys}
    buckets["REMAINING"] = []

    patients_in = {k: set() for k in keys}
    patients_in["REMAINING"] = set()
    images_in   = {k: set() for k in keys}

    def try_place(sample, target_keys):
        pid = patient_id(sample)
        for k in target_keys:
            cap = split_size if k != "REMAINING" else 10**9
            if len(buckets[k]) < cap and sample not in images_in[k] and pid not in patients_in[k]:
                buckets[k].append(sample)
                patients_in[k].add(pid)
                images_in[k].add(sample)
                return True
        for k in target_keys:
            cap = split_size if k != "REMAINING" else 10**9
            if len(buckets[k]) < cap and pid not in patients_in[k]:
                buckets[k].append(sample)
                patients_in[k].add(pid)
                images_in[k].add(sample)
                return True
        return False

    for sample in pool:
        if not try_place(sample, keys):
            buckets["REMAINING"].append(sample)

    return buckets


def split_sort_key(key: str):
    prefix, sep, suffix = key.rpartition("_")
    if sep and suffix.isdigit():
        return prefix, int(suffix)
    return key, 0


def save_splits_to_excel(path: str, splits: dict, annotators: dict):
    ordered = sorted([k for k in splits if k != "REMAINING"], key=split_sort_key)
    if "REMAINING" in splits:
        ordered.append("REMAINING")

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for key in ordered:
            samples = splits[key]
            ann_name = annotators.get(key, "")
            sheet_name = key[:31]
            df = pd.DataFrame({
                "#": list(range(1, len(samples) + 1)),
                "Sample ID": samples,
            })

            df.to_excel(writer, sheet_name=sheet_name, startrow=2, index=False)
            worksheet = writer.sheets[sheet_name]
            worksheet["A1"] = f"Annotator: {ann_name}"


# ── Split Manager App ───────────────────────────────────────────────────────

class SplitApp(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Annotation Split Manager")
        self.resize(1300, 750)

        self._json_data: dict = {}
        self._splits: dict = {}              # key → list of samples
        self._annotators: dict = {}          # key → annotator name
        self._last_excel_path = str(DEFAULT_EXCEL)
        self._save_process = None

        self._load_json()
        self._splits[REPORTX_KEY] = load_reportx_samples()
        self._annotators[REPORTX_KEY] = REPORTX_KEY
        self._build_ui()
        self.setStyleSheet(MAIN_STYLE)
        self._autoload_excel()

    # ── JSON ─────────────────────────────────────────────────────────────────

    def _load_json(self):
        try:
            with open(INPUT_JSON) as f:
                self._json_data = json.load(f)
        except FileNotFoundError:
            QMessageBox.critical(self, "Error", f"Cannot find {INPUT_JSON}")
            sys.exit(1)

    # ── UI ────────────────────────────────────────────────────────────────────

    def _build_ui(self):
        central = QWidget()
        self.setCentralWidget(central)
        root = QHBoxLayout(central)
        root.setSpacing(12)
        root.setContentsMargins(12, 12, 12, 12)

        # ── LEFT PANEL ──────────────────────────────────────────────────────
        left = QWidget()
        left.setFixedWidth(280)
        left_layout = QVBoxLayout(left)
        left_layout.setSpacing(10)
        left_layout.setContentsMargins(0, 0, 0, 0)

        # Configuration
        config_box = QGroupBox("Configuration")
        config_form = QFormLayout(config_box)
        self._train_r = self._spin(1, 1, 10, "Train repeats")
        self._val_r = self._spin(1, 1, 10, "Val repeats")
        self._test_r = self._spin(3, 1, 10, "Test repeats")
        config_form.addRow("Train repeats:", self._train_r)
        config_form.addRow("Val repeats:", self._val_r)
        config_form.addRow("Test repeats:", self._test_r)

        split_box = QGroupBox("Split Configuration")
        split_form = QFormLayout(split_box)
        self._n_clin = self._spin(10, 1, 500, "Number of clinicians")
        self._split_sz = self._spin(60, 1, 9999, "Samples per clinician")
        self._prefix = QLineEdit("BATCH")
        self._prefix.setStyleSheet(LINE_EDIT_STYLE)
        self._seed = self._spin(42, 0, 99999, "Random seed")
        split_form.addRow("# Clinicians:", self._n_clin)
        split_form.addRow("Samples/clinician:", self._split_sz)
        split_form.addRow("Key prefix:", self._prefix)
        split_form.addRow("Random seed:", self._seed)

        gen_btn = QPushButton("Generate Splits")
        gen_btn.setStyleSheet(PRIMARY_BTN_STYLE)
        gen_btn.clicked.connect(self._generate)

        # Load / Save
        io_box = QGroupBox("File Operations")
        io_layout = QVBoxLayout(io_box)
        load_btn = QPushButton("Load from Excel")
        load_btn.setStyleSheet(BUTTON_STYLE)
        load_btn.clicked.connect(self._load_excel)
        self._save_btn = QPushButton("Save to Excel")
        self._save_btn.setStyleSheet(BUTTON_STYLE)
        self._save_btn.clicked.connect(self._save_excel)
        io_layout.addWidget(load_btn)
        io_layout.addWidget(self._save_btn)

        # Clinician Management
        clin_box = QGroupBox("Clinician Management")
        clin_layout = QVBoxLayout(clin_box)

        self._new_btn = QPushButton("+ New Clinician")
        self._new_btn.setStyleSheet(ADD_BTN_STYLE)
        self._new_btn.clicked.connect(self._new_clinician)

        add_layout = QHBoxLayout()
        add_layout.addWidget(QLabel("Add samples:"))
        self._add_sz = self._spin(30, 1, 9999, "Samples to add")
        self._add_btn = QPushButton("+ Add")
        self._add_btn.setStyleSheet(ADD_BTN_STYLE)
        self._add_btn.clicked.connect(self._add_to_clinician)
        add_layout.addWidget(self._add_sz)
        add_layout.addWidget(self._add_btn, 0, Qt.AlignmentFlag.AlignRight)

        self._del_btn = QPushButton("− Delete Clinician")
        self._del_btn.setStyleSheet(DEL_BTN_STYLE)
        self._del_btn.clicked.connect(self._delete_clinician)

        clin_layout.addWidget(self._new_btn)
        clin_layout.addLayout(add_layout)
        clin_layout.addWidget(self._del_btn)

        # Stats
        self._stats_label = QLabel("No splits yet")
        self._stats_label.setWordWrap(True)
        self._stats_label.setStyleSheet(f"color: {SUBTEXT}; font-size: 11px; padding: 8px;")
        self._stats_label.setAlignment(Qt.AlignmentFlag.AlignTop)

        left_layout.addWidget(config_box)
        left_layout.addWidget(split_box)
        left_layout.addWidget(gen_btn)
        left_layout.addWidget(io_box)
        left_layout.addWidget(clin_box)
        left_layout.addWidget(self._stats_label, 1)

        # ── CENTER: TABBED VIEW ─────────────────────────────────────────────
        self._tabs = QTabWidget()
        self._tabs.setStyleSheet(TAB_STYLE)
        self._tabs.tabCloseRequested.connect(lambda: None)

        root.addWidget(left)
        root.addWidget(self._tabs, 1)
        self._update_controls_enabled()

    def _spin(self, val, lo, hi, tooltip=""):
        s = QSpinBox()
        s.setRange(lo, hi)
        s.setValue(val)
        s.setStyleSheet(SPIN_STYLE)
        if tooltip:
            s.setToolTip(tooltip)
        return s

    def _update_controls_enabled(self):
        has_splits = bool(self._splits)
        self._save_btn.setEnabled(has_splits and self._save_process is None)
        self._new_btn.setEnabled(has_splits)
        self._add_sz.setEnabled(has_splits)
        self._add_btn.setEnabled(has_splits)
        self._del_btn.setEnabled(has_splits)

    # ── Generate Splits ─────────────────────────────────────────────────────

    def _generate(self):
        if self._splits:
            reply = QMessageBox.question(
                self,
                "Confirm",
                "Generate new splits and discard the current edits?",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            )
            if reply != QMessageBox.StandardButton.Yes:
                return

        reportx_samples = load_reportx_samples()
        pool = build_pool(
            self._json_data,
            self._train_r.value(),
            self._val_r.value(),
            self._test_r.value(),
            reportx_ids=set(reportx_samples),
        )
        self._splits = assign_to_buckets(
            pool,
            self._n_clin.value(),
            self._split_sz.value(),
            self._prefix.text().strip() or "BATCH",
            self._seed.value(),
        )
        self._annotators = {k: "" for k in self._splits}
        self._splits[REPORTX_KEY] = reportx_samples
        self._annotators[REPORTX_KEY] = REPORTX_KEY
        self._refresh_tabs()
        self._update_stats()
        self._update_controls_enabled()
        QMessageBox.information(self, "Success", "Splits generated!")

    def _refresh_tabs(self):
        self._tabs.clear()
        ordered_keys = sorted(
            [k for k in self._splits if k not in ("REMAINING", REPORTX_KEY)],
            key=split_sort_key,
        )
        if REPORTX_KEY in self._splits:
            ordered_keys.append(REPORTX_KEY)
        if "REMAINING" in self._splits:
            ordered_keys.append("REMAINING")

        for key in ordered_keys:
            self._add_tab(key)

    def _add_tab(self, key: str):
        samples = self._splits.get(key, [])
        table = self._make_table(samples, key)
        label = f"{key}  ({len(samples)})"
        self._tabs.addTab(table, label)

    def _make_table(self, samples: list, key: str) -> QWidget:
        container = QWidget()
        layout = QVBoxLayout(container)
        layout.setContentsMargins(0, 0, 0, 0)

        # Annotator field
        ann_layout = QHBoxLayout()
        ann_layout.addWidget(QLabel("Annotator:"))
        ann_input = QLineEdit(self._annotators.get(key, ""))
        ann_input.setStyleSheet(LINE_EDIT_STYLE)
        ann_input.setMaximumWidth(200)
        ann_input.editingFinished.connect(
            lambda: self._annotators.update({key: ann_input.text()})
        )
        ann_layout.addWidget(ann_input)
        ann_layout.addStretch()

        # Table with edit/delete capabilities
        table = QTableWidget()
        table.setColumnCount(3)
        table.setHorizontalHeaderLabels(["#", "Sample ID", "×"])
        table.setStyleSheet(TABLE_STYLE)
        table.setRowCount(len(samples))

        def delete_row(row_idx, table_key):
            if table_key == REPORTX_KEY:
                return
            if row_idx < 0 or row_idx >= len(self._splits[table_key]):
                return
            sample = self._splits[table_key][row_idx]
            self._splits[table_key].pop(row_idx)
            self._splits["REMAINING"].append(sample)
            self._refresh_tabs()
            self._update_stats()

        for i, sample in enumerate(samples):
            item_num = QTableWidgetItem(str(i + 1))
            item_num.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            item_num.setFlags(item_num.flags() & ~Qt.ItemFlag.ItemIsEditable)
            item_sample = QTableWidgetItem(sample)
            item_sample.setFlags(item_sample.flags() & ~Qt.ItemFlag.ItemIsEditable)
            item_delete = QPushButton("×")
            item_delete.setStyleSheet(f"""
                QPushButton {{
                    background-color: {RED};
                    color: white;
                    border: none;
                    border-radius: 3px;
                    padding: 0px;
                    font-size: 12px;
                    font-weight: bold;
                }}
                QPushButton:hover {{ background-color: #f5a0b5; }}
            """)
            item_delete.setMaximumWidth(30)
            item_delete.clicked.connect(lambda checked, r=i, k=key: delete_row(r, k))

            table.setItem(i, 0, item_num)
            table.setItem(i, 1, item_sample)
            table.setCellWidget(i, 2, item_delete)

        table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.ResizeToContents)
        table.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        table.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeMode.ResizeToContents)
        table.verticalHeader().setVisible(False)
        table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)

        layout.addLayout(ann_layout)
        layout.addWidget(table)
        return container

    def _update_stats(self):
        if not self._splits:
            self._stats_label.setText("No splits generated")
            return
        excluded = {"REMAINING", REPORTX_KEY}
        clin_keys = [k for k in self._splits if k not in excluded]
        total = sum(len(self._splits[k]) for k in clin_keys)
        rem = len(self._splits.get("REMAINING", []))
        reportx_n = len(self._splits.get(REPORTX_KEY, []))
        n_clin = len(clin_keys)
        self._stats_label.setText(
            f"📊 Stats\n\n"
            f"Clinicians: {n_clin}\n"
            f"Total assigned: {total}\n"
            f"REPORTX (fixed): {reportx_n}\n"
            f"Remaining pool: {rem}\n"
            f"Average/clinician: {total // max(n_clin, 1)}"
        )

    # ── Add / Edit / Delete ─────────────────────────────────────────────────

    def _current_key(self) -> str:
        idx = self._tabs.currentIndex()
        if idx < 0:
            return None
        return self._tabs.tabText(idx).split("  ")[0].strip()

    def _can_add_to_clinician(self, sample: str, clinician_key: str) -> bool:
        """Check if sample can be added to clinician (no duplicates, no same-patient)"""
        clinician_samples = self._splits.get(clinician_key, [])
        clinician_patients = {patient_id(s) for s in clinician_samples}

        # No duplicate sample
        if sample in clinician_samples:
            return False
        # No same-patient sample
        if patient_id(sample) in clinician_patients:
            return False
        return True

    def _next_clinician_key(self) -> str:
        prefix = self._prefix.text().strip() or "BATCH"
        used_numbers = []
        for key in self._splits:
            key_prefix, sep, suffix = key.rpartition("_")
            if sep and key_prefix == prefix and suffix.isdigit():
                used_numbers.append(int(suffix))

        next_number = max(used_numbers, default=0) + 1
        while f"{prefix}_{next_number:02d}" in self._splits:
            next_number += 1
        return f"{prefix}_{next_number:02d}"

    def _select_tab(self, key: str):
        for idx in range(self._tabs.count()):
            tab_key = self._tabs.tabText(idx).split("  ")[0].strip()
            if tab_key == key:
                self._tabs.setCurrentIndex(idx)
                return

    def _new_clinician(self):
        if not self._splits:
            self._splits["REMAINING"] = []

        key = self._next_clinician_key()
        self._splits[key] = []
        self._annotators[key] = ""
        self._refresh_tabs()
        self._select_tab(key)
        self._update_stats()
        self._update_controls_enabled()

    def _add_to_clinician(self):
        key = self._current_key()
        if not key:
            QMessageBox.warning(self, "Warning", "Select a clinician tab first")
            return
        if key in ("REMAINING", REPORTX_KEY):
            QMessageBox.warning(self, "Warning", f"Cannot add to {key}")
            return
        rem = self._splits.get("REMAINING", [])
        if not rem:
            QMessageBox.information(self, "Info", "No samples in REMAINING pool")
            return

        # Filter REMAINING to only samples that don't violate constraints
        valid_samples = [s for s in rem if self._can_add_to_clinician(s, key)]
        if not valid_samples:
            QMessageBox.warning(self, "Warning", "No valid samples can be added (all violate patient/duplicate constraints)")
            return

        # Take up to requested size from valid samples
        size = min(self._add_sz.value(), len(valid_samples))
        taken = valid_samples[:size]

        # Remove taken samples from REMAINING and add to clinician
        for sample in taken:
            rem.remove(sample)
        self._splits["REMAINING"] = rem
        self._splits[key].extend(taken)

        self._refresh_tabs()
        self._update_stats()
        self._update_controls_enabled()
        msg = f"Added {size} samples to {key}"
        if size < self._add_sz.value():
            msg += f" (only {size} satisfied constraints)"
        QMessageBox.information(self, "Success", msg)

    def _delete_clinician(self):
        key = self._current_key()
        if not key:
            QMessageBox.warning(self, "Warning", "Select a clinician tab first")
            return
        if key in ("REMAINING", REPORTX_KEY):
            QMessageBox.warning(self, "Warning", f"Cannot delete {key}")
            return
        reply = QMessageBox.question(
            self, "Confirm", f"Delete {key} and return samples to REMAINING?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        )
        if reply != QMessageBox.StandardButton.Yes:
            return
        samples = self._splits.pop(key)
        self._annotators.pop(key, None)
        self._splits["REMAINING"].extend(samples)
        self._refresh_tabs()
        self._update_stats()
        self._update_controls_enabled()

    # ── Save / Load Excel ────────────────────────────────────────────────────

    def _select_save_excel_path(self) -> str:
        dialog = QFileDialog(self, "Save Excel", self._last_excel_path)
        dialog.setAcceptMode(QFileDialog.AcceptMode.AcceptSave)
        dialog.setNameFilter("Excel (*.xlsx)")
        dialog.setDefaultSuffix("xlsx")
        dialog.setOption(QFileDialog.Option.DontUseNativeDialog, True)

        if dialog.exec() != QFileDialog.DialogCode.Accepted:
            return ""
        selected = dialog.selectedFiles()
        return selected[0] if selected else ""

    def _select_load_excel_path(self) -> str:
        dialog = QFileDialog(
            self,
            "Load Excel",
            self._last_excel_path or str(DEFAULT_EXCEL),
        )
        dialog.setAcceptMode(QFileDialog.AcceptMode.AcceptOpen)
        dialog.setFileMode(QFileDialog.FileMode.ExistingFile)
        dialog.setNameFilter("Excel (*.xlsx)")
        dialog.setOption(QFileDialog.Option.DontUseNativeDialog, True)

        if dialog.exec() != QFileDialog.DialogCode.Accepted:
            return ""
        selected = dialog.selectedFiles()
        return selected[0] if selected else ""

    def _save_excel(self):
        if not self._splits:
            QMessageBox.warning(self, "Warning", "No splits to save")
            return

        path = self._select_save_excel_path()
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"
        self._last_excel_path = path

        self._save_btn.setEnabled(False)
        self._save_btn.setText("Saving...")

        payload = {
            "path": path,
            "splits": {k: list(v) for k, v in self._splits.items()},
            "annotators": dict(self._annotators),
        }
        self._save_process = QProcess(self)
        self._update_controls_enabled()
        self._save_btn.setText("Saving...")
        self._save_process.finished.connect(self._on_save_process_finished)
        self._save_process.errorOccurred.connect(self._on_save_process_error)
        self._save_process.start(
            sys.executable,
            [str(Path(__file__).with_name("excel_writer.py"))],
        )
        if self._save_process.waitForStarted(3000):
            self._save_process.write(json.dumps(payload).encode("utf-8"))
            self._save_process.closeWriteChannel()
        else:
            self._finish_save_ui()
            self._cleanup_save_process()
            QMessageBox.critical(self, "Error", "Failed to start Excel writer process")

    def _finish_save_ui(self):
        self._update_controls_enabled()
        self._save_btn.setText("Save to Excel")

    def _cleanup_save_process(self):
        if self._save_process:
            self._save_process.deleteLater()
        self._save_process = None

    def _on_save_process_finished(self, exit_code: int, exit_status):
        self._finish_save_ui()
        process = self._save_process
        stderr = bytes(process.readAllStandardError()).decode(errors="replace") if process else ""
        self._cleanup_save_process()

        if exit_code == 0:
            QMessageBox.information(self, "Success", f"Saved to {Path(self._last_excel_path).name}")
        elif "PermissionError" in stderr:
            QMessageBox.critical(self, "Error", "Permission denied. Is the file open in Excel?")
        else:
            message = stderr.strip() or f"Excel writer exited with code {exit_code}"
            QMessageBox.critical(self, "Error", f"Failed to save: {message}")

    def _on_save_process_error(self, error):
        if error != QProcess.ProcessError.FailedToStart:
            return
        self._finish_save_ui()
        self._cleanup_save_process()
        QMessageBox.critical(self, "Error", "Failed to start Excel writer process")

    def _load_excel(self):
        path = self._select_load_excel_path()
        if not path:
            return

        try:
            self._load_excel_path(path)
            QMessageBox.information(self, "Success", f"Loaded {len(self._splits)} clinicians")
        except Exception as e:
            QMessageBox.critical(self, "Error", f"Failed to load: {str(e)}")

    def _autoload_excel(self):
        if not DEFAULT_EXCEL.exists():
            self._refresh_tabs()
            self._update_stats()
            self._update_controls_enabled()
            return

        try:
            self._load_excel_path(str(DEFAULT_EXCEL))
        except Exception as e:
            self._stats_label.setText(f"Could not autoload {DEFAULT_EXCEL.name}: {e}")
            self._update_controls_enabled()

    def _load_excel_path(self, path: str):
        self._last_excel_path = path
        wb = openpyxl.load_workbook(path)
        self._splits = {}
        self._annotators = {}

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            samples = []
            ann_name = ""

            for cell_ref in ("A1", "A2"):
                if ws[cell_ref].value:
                    ann_text = str(ws[cell_ref].value)
                    if ":" in ann_text:
                        ann_name = ann_text.split(":", 1)[1].strip()
                        break

            sample_start_row = 4
            for row_idx in range(1, ws.max_row + 1):
                if ws[f"B{row_idx}"].value == "Sample ID":
                    sample_start_row = row_idx + 1
                    break

            for row_idx in range(sample_start_row, ws.max_row + 1):
                sample_val = ws[f"B{row_idx}"].value
                if sample_val and isinstance(sample_val, str):
                    samples.append(sample_val)

            self._splits[sheet_name] = samples
            self._annotators[sheet_name] = ann_name

        # REPORTX is always recomputed from the canonical source file,
        # regardless of what was saved in the loaded workbook.
        self._splits[REPORTX_KEY] = load_reportx_samples()
        self._annotators[REPORTX_KEY] = REPORTX_KEY

        self._refresh_tabs()
        self._update_stats()
        self._update_controls_enabled()


# ── Entry point ─────────────────────────────────────────────────────────────

if __name__ == "__main__":
    app = QApplication(sys.argv)
    app.setFont(QFont("Segoe UI", 10) if sys.platform == "win32" else QFont("Ubuntu", 10))
    win = SplitApp()
    win.show()
    sys.exit(app.exec())
